"""Generate the GRN Excel sheet — the demo's final artifact.

Columns mirror what a SAP goods-receipt (MIGO, movement type 101) flat-file needs,
so the sheet is 'SAP-ready': PO + line, material, plant, storage loc, movement
type, UOM, received/accepted/rejected qty, batch/heat, QC result. The dummy PO
supplies material identity; the quality guy supplies qty + QC."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_HEAD = Font(bold=True, color="FFFFFF", size=11)
_HFILL = PatternFill("solid", fgColor="1F4E78")
_TITLE = Font(bold=True, size=15, color="1F4E78")
_LABEL = Font(bold=True, color="333333")
_DEMO = Font(bold=True, size=10, color="B00000")
_thin = Side(style="thin", color="BBBBBB")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def generate_grn(path: str, *, grn_no: str, posting_date: str, po: dict,
                 received_qty: str, accepted_qty: str, rejected_qty: str,
                 qc_result: str, batch_no: str, invoice_no: str | None,
                 remarks: str = "") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "GRN"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "GOODS RECEIPT NOTE (GRN)"
    ws["A1"].font = _TITLE
    ws["A2"] = "DEMO — dummy PO data, stands in for the SAP open-PO feed"
    ws["A2"].font = _DEMO

    meta = [
        ("GRN No.", grn_no), ("Posting Date", posting_date), ("Plant", po["plant"]),
        ("Vendor", po["vendor_name"]), ("Vendor GSTIN", po["vendor_gstin"]),
        ("Vendor Invoice No.", invoice_no or "-"),
    ]
    r = 4
    for label, val in meta:
        ws.cell(r, 1, label).font = _LABEL
        ws.cell(r, 2, val)
        r += 1

    # line-item table
    r += 1
    cols = ["PO No.", "PO Line", "Movement Type", "Material Code", "Material Description",
            "Storage Loc", "UOM", "Received Qty", "Accepted Qty", "Rejected Qty",
            "Batch / Heat", "QC Result"]
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(r, c, name)
        cell.font, cell.fill, cell.border = _HEAD, _HFILL, _BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    vals = [po["po_no"], po["line"], "101", po["material_code"], po["material_desc"],
            po["storage_loc"], po["uom"], received_qty, accepted_qty, rejected_qty,
            batch_no or "-", qc_result.upper()]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(r + 1, c, v)
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center")

    r += 3
    ws.cell(r, 1, "Remarks").font = _LABEL
    ws.cell(r, 2, remarks or "-")
    r += 2
    ws.cell(r, 1, "Quality & Quantity verified by: ____________________").font = _LABEL

    widths = [13, 8, 13, 14, 34, 11, 7, 13, 13, 13, 16, 11]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w

    wb.save(path)
    return path
